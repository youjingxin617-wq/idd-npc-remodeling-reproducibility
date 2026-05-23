from pathlib import Path
import csv
import re

from PIL import Image as PILImage
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image as RLImage, PageBreak, Table, TableStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from pypdf import PdfReader


PROJECT = Path(r"C:\Users\Administrator\Documents\New project 2\idd_scRNA_reproduction")
VAULT = Path(r"E:\github\youwenblood\生信复现sci文献优化")
OUT_DIR = VAULT / "08_BMC_submission"
OUT_DIR.mkdir(parents=True, exist_ok=True)

FIGS = {
    "Figure 1": VAULT / "03_图件总审查/main_figure_candidates/Figure1_global_scRNA_atlas_manual_celltypes_candidate.png",
    "Figure 2": VAULT / "03_图件总审查/main_figure_candidates/Figure2_NPC_subtypes_mild_vs_severe_candidate.png",
    "Figure 3": VAULT / "03_图件总审查/main_figure_candidates/Figure3_ferroptosis_scores_candidate.png",
    "Figure 4": VAULT / "03_图件总审查/main_figure_candidates/Figure4_DEG_enrichment_candidate.png",
    "Figure 5": VAULT / "01_外部bulk验证/figures/Figure5_external_bulk_module_validation_main_candidate.png",
    "Figure 6": VAULT / "05_伪时序分析/figures/56_npc_slingshot_main_Figure6_candidate.png",
    "Figure 7": VAULT / "06_细胞通讯分析/figures/63_npc_cellchat_main_Figure7_candidate.png",
}

TITLE = (
    "Public single-cell and bulk transcriptomic analysis suggests nucleus pulposus cell state "
    "remodeling in intervertebral disc degeneration"
)
RUNNING_TITLE = "Public transcriptomic analysis of NPC remodeling"
KEYWORDS = (
    "Intervertebral disc degeneration; nucleus pulposus cell; single-cell RNA sequencing; "
    "public database; extracellular matrix; ferroptosis-related score"
)
SUBTYPES = [
    "ECM/Adh-NPCs", "Hom-NPCs", "Eff-NPCs", "Ht-NPCs",
    "Fibro-NPCs", "Fibro-reg NPCs", "Hom/Reg-like NPCs"
]

CITATIONS = [
    "Hartvigsen J, Hancock MJ, Kongsted A, Louw Q, Ferreira ML, Genevay S, et al. What low back pain is and why we need to pay attention. Lancet. 2018;391(10137):2356-2367. doi:10.1016/S0140-6736(18)30480-X. PMID:29573870.",
    "Risbud MV, Shapiro IM. Role of cytokines in intervertebral disc degeneration: pain and disc content. Nat Rev Rheumatol. 2014;10(1):44-56. doi:10.1038/nrrheum.2013.160. PMID:24166242.",
    "Tu J, Li W, Yang S, Yang P, Yan Q, Wang S, et al. Single-cell transcriptome profiling reveals multicellular ecosystem of nucleus pulposus during degeneration progression. Adv Sci (Weinh). 2022;9(3):e2103631. doi:10.1002/advs.202103631. PMID:34825784.",
    "Wang Y, Tan L, Yang Y, Duan H, Liu C, Zhao J, et al. Targeting the ROS-ferroptosis-inflammation cycle with a nanozyme-functionalized hydrogel for intervertebral disc repair. Nat Commun. 2025;16:11253. doi:10.1038/s41467-025-66116-w.",
    "Jia S, Liu H, Yang T, Gao S, Li D, Zhang Z, et al. Single-cell sequencing reveals cellular heterogeneity of nucleus pulposus in intervertebral disc degeneration. Sci Rep. 2024;14(1):27245. doi:10.1038/s41598-024-78675-x. PMID:39516278.",
    "Kazezian Z, Gawri R, Haglund L, Ouellet J, Mwale F, Tarrant F, et al. Gene expression profiling identifies interferon signalling molecules and IGFBP3 in human degenerative annulus fibrosus. Sci Rep. 2015;5:15662. doi:10.1038/srep15662. PMID:26489762.",
    "Wang Y, Jiang L, Dai G, Li S, Mu X. Bioinformatics analysis reveals different gene expression patterns in the annulus fibrosis and nucleus pulpous during intervertebral disc degeneration. Exp Ther Med. 2018;16(6):5031-5040. doi:10.3892/etm.2018.6884. PMID:30542457.",
    "Barrett T, Wilhite SE, Ledoux P, Evangelista C, Kim IF, Tomashevsky M, et al. NCBI GEO: archive for functional genomics data sets--update. Nucleic Acids Res. 2013;41(Database issue):D991-D995. doi:10.1093/nar/gks1193. PMID:23193258.",
    "Hao Y, Hao S, Andersen-Nissen E, Mauck WM 3rd, Zheng S, Butler A, et al. Integrated analysis of multimodal single-cell data. Cell. 2021;184(13):3573-3587.e29. doi:10.1016/j.cell.2021.04.048. PMID:34062119.",
    "Korsunsky I, Millard N, Fan J, Slowikowski K, Zhang F, Wei K, et al. Fast, sensitive and accurate integration of single-cell data with Harmony. Nat Methods. 2019;16(12):1289-1296. doi:10.1038/s41592-019-0619-0. PMID:31740819.",
    "Zhou N, Yuan X, Du Q, Zhang Z, Shi X, Bao J, et al. FerrDb V2: update of the manually curated database of ferroptosis regulators and ferroptosis-disease associations. Nucleic Acids Res. 2023;51(D1):D571-D582. doi:10.1093/nar/gkac935. PMID:36305834.",
    "Kuleshov MV, Jones MR, Rouillard AD, Fernandez NF, Duan Q, Wang Z, et al. Enrichr: a comprehensive gene set enrichment analysis web server 2016 update. Nucleic Acids Res. 2016;44(W1):W90-W97. doi:10.1093/nar/gkw377. PMID:27141961.",
    "Street K, Risso D, Fletcher RB, Das D, Ngai J, Yosef N, et al. Slingshot: cell lineage and pseudotime inference for single-cell transcriptomics. BMC Genomics. 2018;19(1):477. doi:10.1186/s12864-018-4772-0. PMID:29914354.",
    "Jin S, Guerrero-Juarez CF, Zhang L, Chang I, Ramos R, Kuan CH, et al. Inference and analysis of cell-cell communication using CellChat. Nat Commun. 2021;12(1):1088. doi:10.1038/s41467-021-21246-9. PMID:33597522.",
    "Jin S, Plikus MV, Nie Q. CellChat for systematic analysis of cell-cell communication from single-cell transcriptomics. Nat Protoc. 2025;20(1):180-219. doi:10.1038/s41596-024-01045-4. PMID:39289562.",
]


def read_csv(path):
    with open(path, newline="", encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


def pct(x):
    return f"{float(x) * 100:.1f}%"


def fmt(x, digits=3):
    try:
        return f"{float(x):.{digits}g}"
    except Exception:
        return str(x)


def esc(text):
    return str(text).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def md_path(path):
    return str(path).replace("\\", "/")


for name, path in FIGS.items():
    if not path.exists():
        raise FileNotFoundError(f"{name}: {path}")

counts_rows = read_csv(PROJECT / "results/tables/npc_final_subtype_counts_and_proportions.csv")
deg_rows = read_csv(PROJECT / "results/tables/clean_npc_deg_summary_by_subtype.csv")
external_rows = read_csv(PROJECT / "results/tables_external_bulk/GSE70362_NP_low_high_module_score_stats.csv")
slingshot = read_csv(PROJECT / "results/tables_slingshot/npc_slingshot_pseudotime_global_mild_vs_severe_stats.csv")[0]
lineages = read_csv(PROJECT / "results/tables_slingshot/npc_slingshot_lineages.csv")
cellchat_summary = read_csv(PROJECT / "results/tables_cellchat/npc_cellchat_global_network_summary.csv")
cellchat_delta = read_csv(PROJECT / "results/tables_cellchat/npc_cellchat_outgoing_incoming_delta_by_subtype.csv")
cellchat_pathways = read_csv(PROJECT / "results/tables_cellchat/npc_cellchat_top_severe_enhanced_pathways.csv")

counts = {s: {"mild_n": "0", "mild_pct": "0.0%", "severe_n": "0", "severe_pct": "0.0%"} for s in SUBTYPES}
for row in counts_rows:
    subtype = row["npc_subtype_final"]
    group = row["article_group"]
    counts[subtype][f"{group}_n"] = row["n_cells"]
    counts[subtype][f"{group}_pct"] = pct(row["proportion"])

deg = {s: {"up": "0", "down": "0"} for s in SUBTYPES}
for row in deg_rows:
    if row["direction"] == "up_in_severe":
        deg[row["npc_subtype_final"]]["up"] = row["n_sig"]
    if row["direction"] == "down_in_severe":
        deg[row["npc_subtype_final"]]["down"] = row["n_sig"]

cc_by_group = {r["group"]: r for r in cellchat_summary}
cc_delta = {r["subtype"]: r for r in cellchat_delta}
top_cc_paths = ", ".join([r["pathway_name"] for r in cellchat_pathways[:8]])

figure_legends = {
    "Figure 1": "Single-cell atlas and NPC-focused workflow for public GSE165722 reanalysis.",
    "Figure 2": "NPC subclustering and degeneration-associated redistribution of NPC states. The Hom/Reg-like NPC label denotes a reclustered homeostatic/regulatory-like state supported by CHRDL2, FRZB, WIF1, IGF2, and CLEC3A.",
    "Figure 3": "FerrDb-derived ferroptosis driver and suppressor scores in NPC subtypes. These are transcriptomic scores, not experimental evidence of ferroptotic cell death.",
    "Figure 4": "Subtype-resolved differential expression and enrichment analysis highlighting matrix remodeling, collagen organization, adhesion signaling, and stress-response programs.",
    "Figure 5": "Exploratory external bulk validation in independent NP transcriptomes. GSE70362 was the main external dataset; GSE147383 was used only as a small supportive trend dataset.",
    "Figure 6": "Slingshot trajectory analysis suggesting a transcriptional continuum from homeostatic/regulatory-like NPC states toward fibro-remodeling NPC states.",
    "Figure 7": "Exploratory CellChat analysis suggesting ECM-centered redistribution of inferred NPC-NPC communication. The main panel summarizes severe-minus-mild communication strength, subtype-level incoming and outgoing communication weights, and severe-enhanced pathways. Full count/weight heatmaps and selected severe-enhanced ligand-receptor pairs are provided in Supplementary Figure S1.",
}

abstract = {
    "Background": "Intervertebral disc degeneration (IDD) involves extracellular matrix disruption and nucleus pulposus (NP) cell dysfunction, but degeneration-associated NP cell states remain incompletely defined.",
    "Methods": "We reanalyzed public human NP single-cell RNA-seq data from GSE165722. After quality control, integration, major cell-type annotation, NPC extraction, and NPC reclustering, cleaned NPC-like cells were compared between mild and severe degeneration. FerrDb-derived ferroptosis driver and suppressor scores, subtype-resolved differential expression, enrichment analysis, exploratory external bulk validation, Slingshot pseudotime, and CellChat communication inference were performed.",
    "Results": "A total of 8,245 cleaned NPC-like cells were analyzed, including 5,809 mild and 2,436 severe cells. Seven NPC states were annotated. Severe degeneration was associated with expansion of Fibro-NPCs (10.8% to 28.9%) and Fibro-reg NPCs (6.8% to 15.4%), and reduction of Hom/Reg-like NPCs (32.4% to 13.4%) and Hom-NPCs (21.1% to 14.9%). Enrichment analysis highlighted extracellular matrix organization, collagen formation, ECM-receptor interaction, focal adhesion, and stress-response programs. In GSE70362, the matrix-homeostasis module was lower and the iron/ferroptosis-response module was higher in high-grade degeneration. Slingshot supported a transcriptional continuum toward fibro-remodeling NPC states, while CellChat suggested communication redistribution rather than global communication enhancement.",
    "Conclusions": "This public transcriptomic reanalysis suggests that IDD severity is associated with NPC state redistribution toward fibro-remodeling and stress-responsive programs. The findings should be interpreted as exploratory and hypothesis-generating, and experimental validation will be required before mechanistic or translational conclusions can be made."
}

sections = [
    ("Introduction", [
        "Intervertebral disc degeneration (IDD) is a major contributor to chronic low back pain and disability [1]. Degenerative remodeling of the intervertebral disc involves extracellular matrix (ECM) disruption, altered collagen and proteoglycan homeostasis, inflammatory and oxidative stress responses, and dysfunction of resident nucleus pulposus (NP) cells [2].",
        "Nucleus pulposus cells (NPCs) maintain the NP matrix microenvironment. During degeneration, NPCs may undergo transcriptional state changes associated with catabolic remodeling, fibrotic matrix deposition, stress adaptation, and cell-death-related programs. Bulk transcriptomic studies can identify tissue-level changes, but they cannot resolve whether these signals arise from distinct NPC states or from changes in cell composition.",
        "Single-cell RNA sequencing provides a framework for dissecting cellular heterogeneity in human disc tissues. The GSE165722 study previously reported a multicellular NP ecosystem during degeneration progression, including fibroNPCs associated with end-stage degeneration [3]. A subsequent therapeutic study used GSE165722 together with experimental models to connect NPC oxidative stress, ferroptosis, inflammation, and hydrogel-based disc repair [4]. A more recent human NP single-cell study also supported substantial cellular heterogeneity during IDD [5].",
        "These studies provide a strong mechanistic and experimental background, but they also leave room for a complementary public-data reanalysis centered on NPC state remodeling rather than therapeutic material development. In particular, degeneration-associated matrix/fibro-remodeling programs, external public bulk transcriptomic support, exploratory state-continuity structure, and inferred NPC-NPC communication changes can be organized as a reproducible secondary transcriptomic resource.",
        "Ferroptosis-related transcriptional programs have attracted attention in degenerative and inflammatory diseases. However, transcriptomic module scoring alone cannot establish that ferroptotic cell death has occurred. In IDD, ferroptosis-related gene sets may nevertheless provide information about iron handling, antioxidant response, and lipid-peroxidation-associated stress programs within specific NPC states.",
        "Here, we performed a focused secondary reanalysis of public human NP single-cell and bulk transcriptomes to characterize degeneration-associated NPC state redistribution, ferroptosis-related transcriptional scores, matrix-remodeling pathways, exploratory external bulk transcriptomic support, pseudotime structure, and inferred NPC-NPC communication remodeling."
    ]),
    ("Materials and Methods", [
        ("Data sources and study design", "This study was a secondary bioinformatics analysis of public transcriptomic datasets. The main single-cell dataset was GSE165722 from the Gene Expression Omnibus [3,8]. External bulk transcriptomic datasets included GSE70362 and GSE147383; GSE70362 has been used in previous disc transcriptomic studies and contains public NP samples with Thompson-grade annotations [6,7]. No new human specimens were collected and no wet-laboratory experiments were performed."),
        ("Single-cell preprocessing and integration", "Raw count matrices from GSE165722 were processed in R using Seurat [9]. Individual sample objects were created with min.cells = 3 and min.features = 200. After sample merging, mitochondrial transcript percentage was calculated using genes matching the ^MT- pattern. Cells were retained when nFeature_RNA > 200, nFeature_RNA < two times the median nFeature_RNA, and percent.mt < 20. Data were log-normalized using Seurat default NormalizeData settings, 2,000 variable features were selected with the vst method, all genes were scaled, and PCA was performed with 30 principal components. Harmony integration was performed using sample_id as the batch variable and PCA dimensions 1-30 [10]. Neighbor graph construction, clustering, UMAP, and t-SNE visualization used the Harmony reduction with dimensions 1-30; the clustering resolution was 0.4."),
        ("NPC extraction and subtype annotation", "Major cell populations were annotated manually according to canonical marker-gene patterns. NPC-like/stromal clusters were extracted for focused analysis and reprocessed using the same normalization, variable-feature selection, scaling, PCA, Harmony integration, clustering, UMAP, and t-SNE settings described above. Obvious non-NPC contaminant clusters were removed before final NPC subtype annotation. The final cleaned NPC subset contained 8,245 NPC-like cells and was annotated into seven NPC states: ECM/Adh-NPCs, Hom-NPCs, Eff-NPCs, Ht-NPCs, Fibro-NPCs, Fibro-reg NPCs, and Hom/Reg-like NPCs. The Hom/Reg-like NPC label was assigned to the reclustered state marked by CHRDL2, FRZB, WIF1, IGF2, and CLEC3A."),
        ("Ferroptosis-related module scoring", "FerrDb-derived ferroptosis driver and suppressor gene sets were filtered to human entries and intersected with genes detected in the cleaned NPC expression matrix [11]. Ferroptosis driver scores and ferroptosis suppressor scores were calculated using Seurat AddModuleScore with the RNA assay and the matched FerrDb driver or suppressor genes as feature sets. Mild and severe cells were compared within NPC subtypes using Wilcoxon rank-sum tests, and P values were adjusted by the Benjamini-Hochberg method. These scores were interpreted as transcriptomic gene-set scores rather than direct evidence of ferroptotic cell death."),
        ("Differential expression and enrichment", "Within each NPC subtype, severe and mild degeneration cells were compared using Seurat FindMarkers with ident.1 = severe, ident.2 = mild, assay = RNA, test.use = wilcox, logfc.threshold = 0, min.pct = 0.10, and only.pos = FALSE [9]. Subtypes were analyzed only when both mild and severe groups contained at least 20 cells. Genes with adjusted p value < 0.05 and absolute average log2 fold change >= 0.25 were considered significant in this reproduction workflow. Up to the top 300 severe-upregulated genes per subtype, ranked by adjusted p value and fold change, were submitted to Enrichr for GO Biological Process 2025, KEGG 2021 Human, and Reactome 2022 enrichment analyses [12]."),
        ("External bulk validation, pseudotime, and communication inference", "Single-cell-derived biological modules were assessed in external bulk NP transcriptomes. In GSE70362, low-grade degeneration was defined as Thompson I/I-II/II NP samples and high-grade degeneration as Thompson IV/V NP samples; Thompson III samples were excluded. Module scores were calculated as the mean of gene-wise z scores for available genes in each module, including matrix-homeostasis, iron/ferroptosis-response, and fibrotic-remodeling modules. Group comparisons used Wilcoxon rank-sum tests with Benjamini-Hochberg adjustment. Slingshot was applied to a SingleCellExperiment object using UMAP coordinates, NPC subtype labels as cluster labels, and Hom/Reg-like NPCs as the starting cluster when present [13]. Pseudotime values from multiple lineages were normalized to 0-1 and averaged into a composite pseudotime score. CellChat was run separately in mild and severe NPCs using the human CellChat database, RNA expression data, group.by = NPC subtype, type = truncatedMean, trim = 0.1, raw.use = TRUE, population.size = TRUE, nboot = 20, seed.use = 20260521, and min.cells = 10 [14,15]. These analyses were treated as hypothesis-generating and were not used to claim patient-level validation or experimental ligand-receptor activity.")
    ]),
    ("Results", [
        ("Single-cell reanalysis supports NPC-focused analysis", "After quality control, integration, clustering, and manual marker-based annotation of GSE165722, a human intervertebral disc cell atlas was generated. Because NPCs are central to NP matrix maintenance and disc degeneration biology, downstream analyses focused on the NPC-like compartment (Figure 1)."),
        ("Severe degeneration is associated with redistribution of NPC states", "Reclustering of the cleaned NPC compartment resolved seven NPC states. Severe degeneration was associated with expansion of Fibro-NPCs from 10.8% in mild NPCs to 28.9% in severe NPCs, and Fibro-reg NPCs from 6.8% to 15.4%. In contrast, Hom/Reg-like NPCs decreased from 32.4% to 13.4%, and Hom-NPCs decreased from 21.1% to 14.9% (Figure 2 and Table 1)."),
        ("Ferroptosis-related scores reveal subtype-specific stress programs", "FerrDb-derived ferroptosis driver and suppressor scores showed subtype-specific alterations. Driver scores differed between mild and severe degeneration in Eff-NPCs, Ht-NPCs, and Fibro-reg NPCs, while suppressor scores were altered in Hom/Reg-like NPCs. These findings indicate heterogeneous ferroptosis-related transcriptional programs rather than uniform ferroptotic cell death (Figure 3)."),
        ("Differential expression and enrichment implicate matrix remodeling", "Subtype-resolved differential expression identified severe degeneration-associated genes in multiple NPC states (Table 2). Enrichment analysis highlighted ECM-receptor interaction and focal adhesion in Eff-NPCs, extracellular matrix organization and collagen biosynthesis/formation in Fibro-reg NPCs, and glycosaminoglycan or keratan sulfate-related processes in ECM/Adh-NPCs (Figure 4)."),
        ("External bulk transcriptomes provide exploratory support", "In GSE70362 NP samples, the matrix-homeostasis module was lower in high-grade degeneration than in low-grade degeneration (mean score, -0.489 vs 0.612; adjusted P = 0.0218). The iron/ferroptosis-response module was higher in high-grade degeneration (mean score, 0.402 vs -0.503; adjusted P = 0.0218). The fibrotic-remodeling module showed a non-significant upward trend (mean score, 0.246 vs -0.308; adjusted P = 0.625). GSE147383 was treated only as supportive trend evidence because of very small NP sample size (Figure 5 and Table 3)."),
        ("Slingshot supports a continuum toward fibro-remodeling NPC states", f"Slingshot inferred two main NPC trajectories: {lineages[0]['clusters']} and {lineages[1]['clusters']}. Severe NPCs were positioned at later pseudotime values than mild NPCs at the descriptive cell level (mean pseudotime, {float(slingshot['mean_severe']):.3f} vs {float(slingshot['mean_mild']):.3f}; severe-minus-mild delta, {float(slingshot['delta_severe_minus_mild']):.3f}). Slingshot pseudotime correlated strongly with the preceding MST-based pseudotime (Spearman correlation, {float(slingshot['mst_correlation']):.3f}) (Figure 6)."),
        ("CellChat suggests communication redistribution", f"CellChat did not indicate a global increase in NPC-NPC communication in severe degeneration. Total inferred interaction count was lower in severe NPCs than in mild NPCs ({cc_by_group['severe']['total_interaction_count']} vs {cc_by_group['mild']['total_interaction_count']}), and total interaction weight was also slightly lower ({float(cc_by_group['severe']['total_interaction_weight']):.3f} vs {float(cc_by_group['mild']['total_interaction_weight']):.3f}). However, Fibro-NPCs showed increased outgoing communication weight and Fibro-reg NPCs showed increases in both outgoing and incoming weights. Severe-enhanced pathways included {top_cc_paths}. These findings suggest ECM-centered communication remodeling rather than global communication activation (Figure 7).")
    ]),
    ("Discussion", [
        "This public transcriptomic reanalysis suggests that IDD severity is associated with coordinated NPC state redistribution, matrix-remodeling transcriptional programs, subtype-specific ferroptosis-related scores, and remodeling of inferred NPC-NPC communication. The strongest finding is the expansion of Fibro-NPCs and Fibro-reg NPCs in severe degeneration, accompanied by reduced Hom/Reg-like NPC and Hom-NPC proportions.",
        "The enrichment of extracellular matrix organization, collagen formation, ECM-receptor interaction, and focal adhesion pathways provides biological coherence for the Fibro-NPC/Fibro-reg NPC expansion. These states may reflect a broader degeneration-associated fibro-remodeling program involving collagen deposition, altered matrix-cell adhesion, and changes in NP tissue architecture.",
        "Ferroptosis-related module scoring adds a stress-response dimension to the NPC remodeling story. The observed score changes were not uniform across all NPCs, suggesting that ferroptosis-related transcriptional programs are state-dependent. Experimental assays would be required to establish lipid peroxidation, iron dependence, or ferroptotic vulnerability in specific NPC subtypes.",
        "The Slingshot trajectory analysis supports a transcriptional continuum from homeostatic/regulatory-like NPC states toward fibro-remodeling states. This result is consistent with, but does not prove, a degeneration-associated state transition. Pseudotime should be interpreted as a model of transcriptional continuity, not as direct temporal progression, causality, or lineage conversion.",
        "CellChat analysis suggests that severe degeneration involves communication redistribution rather than global communication activation. Fibro-NPCs and Fibro-reg NPCs became more prominent in selected communication axes, especially collagen-, fibronectin/integrin-, THBS-, tenascin-, and periostin-related pathways. These inferred interactions require protein-level or functional validation.",
        "This study has limitations. All analyses are based on public datasets and are constrained by the original sample collection, metadata quality, sequencing depth, and batch structure. Cell-level statistical comparisons should not be mistaken for patient-level replication. The present analysis also reuses a public dataset that has been investigated in previous single-cell and therapeutic studies, so its value lies in focused reanalysis and transparent integration rather than independent experimental discovery. External bulk validation was exploratory, and GSE147383 was limited by very small NP sample size. Ferroptosis scores, pseudotime, and CellChat outputs are computational inferences rather than experimental validation."
    ]),
    ("Conclusion", [
        "Public single-cell and bulk transcriptomic analyses suggest that IDD severity is associated with redistribution of NPC states and enrichment of fibro-remodeling and stress-responsive transcriptional programs. This work provides a reproducible resource for hypothesis generation, but does not establish validated biomarkers, causal mechanisms, or therapeutic targets."
    ])
]

tables = {
    "Table 1": {
        "title": "NPC subtype composition in mild and severe degeneration.",
        "headers": ["NPC subtype", "Mild n", "Mild %", "Severe n", "Severe %"],
        "rows": [[s, counts[s]["mild_n"], counts[s]["mild_pct"], counts[s]["severe_n"], counts[s]["severe_pct"]] for s in SUBTYPES],
        "note": "Percentages were calculated within mild or severe NPCs in the cleaned NPC subset."
    },
    "Table 2": {
        "title": "Subtype-resolved significant differentially expressed genes in severe versus mild degeneration.",
        "headers": ["NPC subtype", "Up in severe", "Down in severe"],
        "rows": [[s, deg[s]["up"], deg[s]["down"]] for s in SUBTYPES],
        "note": "Threshold: adjusted p value < 0.05 and absolute average log2 fold change >= 0.25."
    },
    "Table 3": {
        "title": "Exploratory external validation of module scores in GSE70362 NP samples.",
        "headers": ["Module", "n low", "n high", "Mean low", "Mean high", "Adjusted P"],
        "rows": [[r["module"], r["n_low"], r["n_high"], fmt(r["mean_low"]), fmt(r["mean_high"]), fmt(r["p_adj"])] for r in external_rows],
        "note": "Low-grade group: Thompson I/I-II/II NP samples, n = 8. High-grade group: Thompson IV/V NP samples, n = 10. Thompson III samples were excluded."
    }
}

DECLARATIONS = [
    ("Ethics approval and consent to participate", "Not applicable. This study reanalyzed publicly available transcriptomic datasets and did not involve newly collected human specimens, direct patient contact, or new animal experiments."),
    ("Consent for publication", "Not applicable."),
    ("Availability of data and materials", "The datasets analyzed in this study are publicly available from the Gene Expression Omnibus, including GSE165722, GSE70362, and GSE147383. FerrDb-derived ferroptosis-related gene sets were obtained from FerrDb. Analysis scripts, processed result tables, supplementary tables, and software/session information are available at https://github.com/youjingxin617-wq/idd-npc-remodeling-reproducibility."),
    ("Competing interests", "The authors declare that they have no competing interests. This statement should be confirmed by all authors before submission."),
    ("Funding", "No specific funding information has been provided at this draft stage. This section should be updated after discussion with the supervisor and affiliated department."),
    ("Authors' contributions", "Author contributions are placeholders at this stage. The final version must reflect the actual contributions of all authors according to journal and institutional requirements."),
    ("Acknowledgements", "Not applicable at this draft stage."),
]


def build_markdown(path):
    lines = [f"# {TITLE}", "", f"Running title: {RUNNING_TITLE}", "", "BMC Musculoskeletal Disorders-oriented draft", "", "## Abstract", ""]
    for key, value in abstract.items():
        lines += [f"**{key}:** {value}", ""]
    lines += [f"**Keywords:** {KEYWORDS}", ""]
    for heading, content in sections:
        lines += [f"## {heading}", ""]
        for item in content:
            if isinstance(item, tuple):
                lines += [f"### {item[0]}", "", item[1], ""]
            else:
                lines += [item, ""]
    lines += ["## Tables", ""]
    for tname, table in tables.items():
        lines += [f"### {tname}. {table['title']}", ""]
        lines += ["| " + " | ".join(table["headers"]) + " |"]
        lines += ["|" + "|".join(["---"] * len(table["headers"])) + "|"]
        for row in table["rows"]:
            lines += ["| " + " | ".join(row) + " |"]
        lines += ["", f"Note: {table['note']}", ""]
    lines += ["## Figures", ""]
    for name, fig in FIGS.items():
        lines += [f"![{name}]({md_path(fig)})", "", f"**{name}.** {figure_legends[name]}", ""]
    lines += ["## Declarations", ""]
    for heading, text in DECLARATIONS:
        lines += [f"### {heading}", "", text, ""]
    lines += ["## References", ""]
    for i, ref in enumerate(CITATIONS, 1):
        lines += [f"{i}. {ref}"]
    path.write_text("\n".join(lines), encoding="utf-8")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_doc_picture(doc, path):
    doc.add_picture(str(path), width=Inches(6.35))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_doc_table(doc, tname, table):
    p = doc.add_paragraph()
    p.add_run(f"{tname}. {table['title']}").bold = True
    tb = doc.add_table(rows=1, cols=len(table["headers"]))
    tb.style = "Table Grid"
    tb.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(table["headers"]):
        tb.rows[0].cells[i].text = h
        set_cell_shading(tb.rows[0].cells[i], "D9EAF7")
    for row in table["rows"]:
        cells = tb.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = val
    note = doc.add_paragraph(f"Note: {table['note']}")
    note.runs[0].italic = True


def build_docx(path):
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.65)
    sec.bottom_margin = Inches(0.65)
    sec.left_margin = Inches(0.75)
    sec.right_margin = Inches(0.75)
    for style in doc.styles:
        if style.type == 1:
            style.font.name = "Times New Roman"
            style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
            if style.name == "Normal":
                style.font.size = Pt(10.5)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(TITLE)
    run.bold = True
    run.font.size = Pt(16)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run(f"Running title: {RUNNING_TITLE}").italic = True
    sub2 = doc.add_paragraph()
    sub2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub2.add_run("BMC Musculoskeletal Disorders-oriented draft").italic = True

    doc.add_heading("Abstract", level=1)
    for key, value in abstract.items():
        p = doc.add_paragraph()
        p.add_run(key + ": ").bold = True
        p.add_run(value)
    p = doc.add_paragraph()
    p.add_run("Keywords: ").bold = True
    p.add_run(KEYWORDS)

    for heading, content in sections:
        doc.add_heading(heading, level=1)
        for item in content:
            text = item[1] if isinstance(item, tuple) else item
            if isinstance(item, tuple):
                doc.add_heading(item[0], level=2)
            doc.add_paragraph(text)
            if heading == "Results":
                for fig_name in FIGS:
                    if fig_name in text:
                        add_doc_picture(doc, FIGS[fig_name])
                        cap = doc.add_paragraph()
                        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        cap.add_run(fig_name + ". ").bold = True
                        cap.add_run(figure_legends[fig_name])
                        break

    doc.add_heading("Tables", level=1)
    for tname, table in tables.items():
        add_doc_table(doc, tname, table)

    doc.add_heading("Figure Legends", level=1)
    for name in FIGS:
        p = doc.add_paragraph()
        p.add_run(name + ". ").bold = True
        p.add_run(figure_legends[name])

    doc.add_heading("Declarations", level=1)
    for heading, text in DECLARATIONS:
        doc.add_heading(heading, level=2)
        doc.add_paragraph(text)

    doc.add_heading("References", level=1)
    for i, ref in enumerate(CITATIONS, 1):
        doc.add_paragraph(f"{i}. {ref}")
    doc.save(path)


def register_pdf_fonts():
    try:
        pdfmetrics.registerFont(TTFont("MSYH", r"C:\Windows\Fonts\msyh.ttc"))
        pdfmetrics.registerFont(TTFont("MSYH-Bold", r"C:\Windows\Fonts\msyhbd.ttc"))
        return "MSYH", "MSYH-Bold"
    except Exception:
        return "Times-Roman", "Times-Bold"


def pdf_styles():
    regular, bold = register_pdf_fonts()
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle("TitleX", fontName=bold, fontSize=15, leading=20, alignment=TA_CENTER, spaceAfter=8))
    styles.add(ParagraphStyle("SubX", fontName=regular, fontSize=9, leading=12, alignment=TA_CENTER, textColor=colors.darkgrey, spaceAfter=10))
    styles.add(ParagraphStyle("Head1X", fontName=bold, fontSize=13, leading=17, spaceBefore=12, spaceAfter=6))
    styles.add(ParagraphStyle("Head2X", fontName=bold, fontSize=10.5, leading=14, spaceBefore=7, spaceAfter=4))
    styles.add(ParagraphStyle("BodyX", fontName=regular, fontSize=9, leading=13, alignment=TA_JUSTIFY, spaceAfter=6))
    styles.add(ParagraphStyle("CapX", fontName=regular, fontSize=8, leading=11, alignment=TA_CENTER, textColor=colors.darkgrey, spaceAfter=8))
    styles.add(ParagraphStyle("SmallX", fontName=regular, fontSize=8, leading=11, textColor=colors.darkgrey, spaceAfter=5))
    return styles, regular


def add_pdf_image(story, path, max_w=17.2 * cm, max_h=12.0 * cm):
    img = PILImage.open(path)
    w, h = img.size
    ratio = min(max_w / w, max_h / h)
    story.append(RLImage(str(path), width=w * ratio, height=h * ratio))


def add_pdf_table(story, styles, font, tname, table):
    story.append(Paragraph(f"<b>{esc(tname)}. {esc(table['title'])}</b>", styles["BodyX"]))
    data = [table["headers"]] + table["rows"]
    tb = Table(data, repeatRows=1, hAlign="CENTER")
    tb.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), font),
        ("FONTSIZE", (0, 0), (-1, -1), 7.3),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#D9EAF7")),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (1, 1), (-1, -1), "CENTER"),
    ]))
    story.append(tb)
    story.append(Paragraph(f"Note: {esc(table['note'])}", styles["SmallX"]))


def build_pdf(path):
    styles, font = pdf_styles()
    story = [
        Paragraph(esc(TITLE), styles["TitleX"]),
        Paragraph(esc(f"Running title: {RUNNING_TITLE}"), styles["SubX"]),
        Paragraph("BMC Musculoskeletal Disorders-oriented draft", styles["SubX"]),
        Paragraph("Abstract", styles["Head1X"]),
    ]
    for key, value in abstract.items():
        story.append(Paragraph(f"<b>{esc(key)}:</b> {esc(value)}", styles["BodyX"]))
    story.append(Paragraph(f"<b>Keywords:</b> {esc(KEYWORDS)}", styles["BodyX"]))

    for heading, content in sections:
        story.append(Paragraph(esc(heading), styles["Head1X"]))
        for item in content:
            text = item[1] if isinstance(item, tuple) else item
            if isinstance(item, tuple):
                story.append(Paragraph(esc(item[0]), styles["Head2X"]))
            story.append(Paragraph(esc(text), styles["BodyX"]))
            if heading == "Results":
                for fig_name in FIGS:
                    if fig_name in text:
                        add_pdf_image(story, FIGS[fig_name])
                        story.append(Paragraph(f"<b>{fig_name}.</b> {esc(figure_legends[fig_name])}", styles["CapX"]))
                        break

    story.append(PageBreak())
    story.append(Paragraph("Tables", styles["Head1X"]))
    for tname, table in tables.items():
        add_pdf_table(story, styles, font, tname, table)
        story.append(Spacer(1, 6))

    story.append(PageBreak())
    story.append(Paragraph("Figure Legends", styles["Head1X"]))
    for name in FIGS:
        story.append(Paragraph(f"<b>{esc(name)}.</b> {esc(figure_legends[name])}", styles["BodyX"]))
    story.append(Paragraph("Declarations", styles["Head1X"]))
    for heading, text in DECLARATIONS:
        story.append(Paragraph(esc(heading), styles["Head2X"]))
        story.append(Paragraph(esc(text), styles["BodyX"]))
    story.append(Paragraph("References", styles["Head1X"]))
    for i, ref in enumerate(CITATIONS, 1):
        story.append(Paragraph(f"{i}. {esc(ref)}", styles["BodyX"]))

    doc = SimpleDocTemplate(str(path), pagesize=A4, rightMargin=1.5 * cm, leftMargin=1.5 * cm, topMargin=1.4 * cm, bottomMargin=1.4 * cm)
    doc.build(story)


md_out = OUT_DIR / "BMC_Musculoskeletal_Disorders_oriented_manuscript.md"
docx_out = OUT_DIR / "BMC_Musculoskeletal_Disorders_oriented_manuscript.docx"
pdf_out = OUT_DIR / "BMC_Musculoskeletal_Disorders_oriented_manuscript.pdf"

build_markdown(md_out)
build_docx(docx_out)
build_pdf(pdf_out)

reader = PdfReader(str(pdf_out))
text = "\n".join(page.extract_text() or "" for page in reader.pages[:2])
print("WROTE", md_out)
print("WROTE", docx_out)
print("WROTE", pdf_out)
print("PDF_PAGES", len(reader.pages))
print("PDF_TEXT_OK", bool(re.search("Public single-cell", text)))


